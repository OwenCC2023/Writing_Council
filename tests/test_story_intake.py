import pytest
from docx import Document

from story_intake import load_story_text, word_count


def test_loads_txt(tmp_path):
    p = tmp_path / "story.txt"
    p.write_text("Once upon a time.\n\nThe end.", encoding="utf-8")
    assert load_story_text(p) == "Once upon a time.\n\nThe end."


def test_loads_md(tmp_path):
    p = tmp_path / "story.md"
    p.write_text("# Title\n\nProse here.", encoding="utf-8")
    assert "Prose here." in load_story_text(p)


def test_loads_docx_joining_paragraphs_with_blank_lines(tmp_path):
    p = tmp_path / "story.docx"
    doc = Document()
    doc.add_paragraph("Owen Cardwell-Copenhefer")
    doc.add_paragraph("First line of narrative.")
    doc.save(p)
    text = load_story_text(p)
    assert "Owen Cardwell-Copenhefer\n\nFirst line of narrative." in text


def test_docx_front_matter_is_not_stripped(tmp_path):
    """Stripping is the intake prompt's job, not this function's."""
    p = tmp_path / "story.docx"
    doc = Document()
    doc.add_paragraph("approx. 8,000 words")
    doc.add_paragraph("Narrative starts.")
    doc.save(p)
    assert "approx. 8,000 words" in load_story_text(p)


def test_rejects_unsupported_extension(tmp_path):
    p = tmp_path / "story.pdf"
    p.write_bytes(b"%PDF-1.4")
    with pytest.raises(ValueError, match="Unsupported"):
        load_story_text(p)


def test_rejects_empty_file(tmp_path):
    p = tmp_path / "story.txt"
    p.write_text("   \n\n  ", encoding="utf-8")
    with pytest.raises(ValueError, match="empty"):
        load_story_text(p)


def test_word_count_counts_whitespace_separated_tokens():
    assert word_count("one two  three\nfour") == 4
    assert word_count("") == 0
